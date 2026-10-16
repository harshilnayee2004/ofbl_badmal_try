import os
import io
import zipfile
from fpdf import FPDF
from openpyxl import Workbook
from models.box_model import Box
from models.lot_model import Lot
from fastapi import HTTPException
from models.user_model import User
from models.fuze_model import Fuze
from models.primer_model import Primer
from sqlmodel import Session, select
from models.box_log_model import BoxLog
from models.fuze_log_model import FuzeLog
from models.box_h2a_inspection_model import BoxH2AInspection
from models.box_G70_inspection_model import BoxG70Inspection
from models.box_M104_inspection_model import BoxM104Inspection
from models.primer_1al_inspection_model import Primer1ALInspection
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

class ReportService:
    @staticmethod
    def _build_box_excel_sheet(ws, box, lot, user_info, inspection, box_logs):
        # Styles
        header_font = Font(bold=True, size=14, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        subheader_font = Font(bold=True, size=12)
        subheader_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        label_font = Font(bold=True)
        thin_border = Border(left=Side(style='thin'), 
                     right=Side(style='thin'), 
                     top=Side(style='thin'), 
                     bottom=Side(style='thin'))

        def create_header(row, text, col_span=2):
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
            return row + 1
            
        def create_subheader(row, text, col_span=2):
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = subheader_font
            cell.fill = subheader_fill
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
            return row + 1

        def create_data_row(row, label, val):
            l_cell = ws.cell(row=row, column=1, value=label)
            l_cell.font = label_font
            l_cell.border = thin_border
            v_cell = ws.cell(row=row, column=2, value=str(val) if val is not None else "N/A")
            v_cell.border = thin_border
            return row + 1

        current_row = 1
        
        # Title
        current_row = create_header(current_row, f"Inspection Report: Box {box.box_number or 'Unnumbered'}")
        current_row += 1

        # 1. Lot Details
        current_row = create_subheader(current_row, "1. Lot Details")
        current_row = create_data_row(current_row, "Lot ID", lot.id)
        current_row = create_data_row(current_row, "Lot Number", lot.lot_number)
        current_row = create_data_row(current_row, "Lot Inspection Status", lot.inspection_status)
        current_row = create_data_row(current_row, "Box Type", lot.box_type)
        current_row = create_data_row(current_row, "Total Boxes", lot.number_of_box)
        current_row = create_data_row(current_row, "Total Fuzes", lot.number_of_fuze)
        current_row = create_data_row(current_row, "Fuze Type", lot.fuze_type)
        current_row += 1

        # 2. Box Details
        current_row = create_subheader(current_row, "2. Box Details")
        current_row = create_data_row(current_row, "Box ID", box.id)
        current_row = create_data_row(current_row, "Box Number", box.box_number)
        current_row = create_data_row(current_row, "Sequence Number", box.sequence_number)
        current_row = create_data_row(current_row, "Box Type", box.box_type)
        current_row = create_data_row(current_row, "Inspection Status", box.box_inspection_status)
        current_row = create_data_row(current_row, "Acceptance Status", "Passed" if box.acceptance_status is True else ("Rejected" if box.acceptance_status is False else "Pending/Not Answered"))
        current_row = create_data_row(current_row, "Inspecting User", user_info)
        current_row = create_data_row(current_row, "Start Time", box.start_time.strftime("%Y-%m-%d %H:%M:%S") if box.start_time else "N/A")
        current_row = create_data_row(current_row, "End Time", box.end_time.strftime("%Y-%m-%d %H:%M:%S") if box.end_time else "N/A")
        current_row += 1
        
        # 3. Check List
        current_row = create_subheader(current_row, "3. Inspection Checklist")
        if inspection:
            for column in inspection.__table__.columns:
                if column.name not in ["id"]:
                    val = getattr(inspection, column.name)
                    pretty_key = column.name.replace("_", " ").title()
                    status_text = "N/A"
                    if val is True:
                        status_text = "Pass"
                    elif val is False:
                        status_text = "Fail"
                    current_row = create_data_row(current_row, pretty_key, status_text)
        else:
            current_row = create_data_row(current_row, "Checklist Status", "No checklist found")
        current_row += 1
        
        # 4. Logs
        current_row = create_subheader(current_row, "4. Box Logs", col_span=3)
        
        # Log Headers
        h1 = ws.cell(row=current_row, column=1, value="Timestamp")
        h1.font = label_font
        h1.border = thin_border
        
        h2 = ws.cell(row=current_row, column=2, value="Description")
        h2.font = label_font
        h2.border = thin_border
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=3)
        current_row += 1
        
        if box_logs:
            for log in box_logs:
                l1 = ws.cell(row=current_row, column=1, value=log.date_time.strftime("%Y-%m-%d %H:%M:%S") if log.date_time else "N/A")
                l1.border = thin_border
                
                l2 = ws.cell(row=current_row, column=2, value=log.description)
                l2.border = thin_border
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=3)
                
                # add border to the merged end cell explicitly
                l3 = ws.cell(row=current_row, column=3)
                l3.border = thin_border
                
                current_row += 1
        else:
            l = ws.cell(row=current_row, column=1, value="No logs found")
            l.border = thin_border
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
            current_row += 1

        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 30

    @staticmethod
    def _build_fuze_excel_sheet(ws, fuze, lot, user_info, inspection, fuze_logs):
        # Styles
        header_font = Font(bold=True, size=14, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        subheader_font = Font(bold=True, size=12)
        subheader_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        label_font = Font(bold=True)
        thin_border = Border(left=Side(style='thin'), 
                    right=Side(style='thin'), 
                    top=Side(style='thin'), 
                    bottom=Side(style='thin'))

        def create_header(row, text, col_span=2):
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
            return row + 1
            
        def create_subheader(row, text, col_span=2):
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = subheader_font
            cell.fill = subheader_fill
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
            return row + 1

        def create_data_row(row, label, val):
            l_cell = ws.cell(row=row, column=1, value=label)
            l_cell.font = label_font
            l_cell.border = thin_border
            v_cell = ws.cell(row=row, column=2, value=str(val) if val is not None else "N/A")
            v_cell.border = thin_border
            return row + 1

        current_row = 1
        
        # Title
        current_row = create_header(current_row, f"Inspection Report: Fuze Sequence {fuze.sequence_number}")
        current_row += 1

        # 1. Lot Details
        current_row = create_subheader(current_row, "1. Lot Details")
        current_row = create_data_row(current_row, "Lot ID", lot.id)
        current_row = create_data_row(current_row, "Lot Number", lot.lot_number)
        current_row = create_data_row(current_row, "Lot Inspection Status", lot.inspection_status)
        current_row = create_data_row(current_row, "Box Type", lot.box_type)
        current_row = create_data_row(current_row, "Total Boxes", lot.number_of_box)
        current_row = create_data_row(current_row, "Total Fuzes", lot.number_of_fuze)
        current_row = create_data_row(current_row, "Fuze Type", lot.fuze_type)
        current_row += 1

        # 2. Fuze Details
        current_row = create_subheader(current_row, "2. Fuze Details")
        current_row = create_data_row(current_row, "Fuze ID", fuze.id)
        current_row = create_data_row(current_row, "Sequence Number", fuze.sequence_number)
        current_row = create_data_row(current_row, "Fuze Type", fuze.fuze_type)
        current_row = create_data_row(current_row, "Inspection Status", fuze.fuze_inspection_status)
        current_row = create_data_row(current_row, "Acceptance Status", "Passed" if fuze.acceptance_status is True else ("Rejected" if fuze.acceptance_status is False else "Pending/Not Answered"))
        current_row = create_data_row(current_row, "Inspecting User", user_info)
        current_row = create_data_row(current_row, "Start Time", fuze.start_time.strftime("%Y-%m-%d %H:%M:%S") if fuze.start_time else "N/A")
        current_row = create_data_row(current_row, "End Time", fuze.end_time.strftime("%Y-%m-%d %H:%M:%S") if fuze.end_time else "N/A")
        current_row += 1
        
        # 3. Check List
        current_row = create_subheader(current_row, "3. Inspection Checklist")
        if inspection:
            for column in inspection.__table__.columns:
                if column.name not in ["id"]:
                    val = getattr(inspection, column.name)
                    pretty_key = column.name.replace("_", " ").title()
                    status_text = "N/A"
                    if val is True:
                        status_text = "Pass"
                    elif val is False:
                        status_text = "Fail"
                    current_row = create_data_row(current_row, pretty_key, status_text)
        else:
            current_row = create_data_row(current_row, "Checklist Status", "No checklist found")
        current_row += 1
        
        # 4. Logs
        current_row = create_subheader(current_row, "4. Fuze Logs", col_span=3)
        
        # Log Headers
        h1 = ws.cell(row=current_row, column=1, value="Timestamp")
        h1.font = label_font
        h1.border = thin_border
        
        h2 = ws.cell(row=current_row, column=2, value="Description")
        h2.font = label_font
        h2.border = thin_border
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=3)
        current_row += 1
        
        if fuze_logs:
            for log in fuze_logs:
                l1 = ws.cell(row=current_row, column=1, value=log.date_time.strftime("%Y-%m-%d %H:%M:%S") if log.date_time else "N/A")
                l1.border = thin_border
                
                l2 = ws.cell(row=current_row, column=2, value=log.description)
                l2.border = thin_border
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=3)
                
                # add border to the merged end cell explicitly
                l3 = ws.cell(row=current_row, column=3)
                l3.border = thin_border
                
                current_row += 1
        else:
            l = ws.cell(row=current_row, column=1, value="No logs found")
            l.border = thin_border
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
            current_row += 1

        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 30

    @staticmethod
    async def generate_box_excel_report(db: Session, box_id: int):
        box = db.exec(select(Box).where(Box.id == box_id)).first()
        if not box:
            raise HTTPException(status_code=404, detail="Box not found")
            
        lot = db.exec(select(Lot).where(Lot.id == box.lot_id)).first()
        if not lot:
            raise HTTPException(status_code=404, detail="Associated lot not found")
        
        box_logs = db.exec(select(BoxLog).where(BoxLog.box_id == box_id).order_by(BoxLog.date_time)).all()
        
        # Determine inspection
        inspection = None
        if getattr(box, "box_h2a_inspection_id", None):
            inspection = db.exec(
                select(BoxH2AInspection).where(BoxH2AInspection.id == box.box_h2a_inspection_id)).first()
        elif getattr(box, "box_type", None) in ["H2A", "Box H2A", "H2A Box"]:
            if getattr(box, "box_h2a_inspection_id", None):
                inspection = db.exec(
                    select(BoxH2AInspection).where(BoxH2AInspection.id == box.box_h2a_inspection_id)).first()
        elif getattr(box, "box_type", None) in ["M-104 Type 2", "M104"] and hasattr(box, "box_m104_inspection_id") and box.box_m104_inspection_id:
            inspection = db.exec(
                select(BoxM104Inspection).where(BoxM104Inspection.id == box.box_m104_inspection_id)).first()
        elif hasattr(box, "box_g70_inspection_id") and box.box_g70_inspection_id:
            inspection = db.exec(
                select(BoxG70Inspection).where(BoxG70Inspection.id == box.box_g70_inspection_id)).first()
        
        user_info = "N/A"
        if box.user_id:
             user = db.exec(select(User).where(User.id == box.user_id)).first()
             if user:
                 user_info = f"({user.email})"

        # Initialize workbook
        wb = Workbook()
        ws = wb.active
        ws.title = f"Box {box.box_number or 'Unnumbered'} Report"

        # Call helper
        ReportService._build_box_excel_sheet(ws, box, lot, user_info, inspection, box_logs)
        
        # Save to memory
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output

    @staticmethod
    async def generate_box_pdf_report(db: Session, box_id: int):
        box = db.exec(select(Box).where(Box.id == box_id)).first()
        if not box:
            raise HTTPException(status_code=404, detail="Box not found")
            
        lot = db.exec(select(Lot).where(Lot.id == box.lot_id)).first()
        if not lot:
            raise HTTPException(status_code=404, detail="Associated lot not found")
        
        box_logs = db.exec(select(BoxLog).where(BoxLog.box_id == box_id).order_by(BoxLog.date_time)).all()
        
        # Determine inspection
        inspection = None
        if getattr(box, "box_h2a_inspection_id", None):
            inspection = db.exec(
                select(BoxH2AInspection).where(BoxH2AInspection.id == box.box_h2a_inspection_id)).first()
        elif getattr(box, "box_type", None) in ["H2A", "Box H2A", "H2A Box"]:
            if getattr(box, "box_h2a_inspection_id", None):
                inspection = db.exec(
                    select(BoxH2AInspection).where(BoxH2AInspection.id == box.box_h2a_inspection_id)).first()
        elif getattr(box, "box_type", None) in ["M-104 Type 2", "M104"] and hasattr(box, "box_m104_inspection_id") and box.box_m104_inspection_id:
            inspection = db.exec(
                select(BoxM104Inspection).where(BoxM104Inspection.id == box.box_m104_inspection_id)).first()
        elif hasattr(box, "box_g70_inspection_id") and box.box_g70_inspection_id:
            inspection = db.exec(
                select(BoxG70Inspection).where(BoxG70Inspection.id == box.box_g70_inspection_id)).first()
        
        user_info = "N/A"
        if box.user_id:
             user = db.exec(select(User).where(User.id == box.user_id)).first()
             if user:
                 user_info = f"({user.email})"

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("helvetica", size=10)
        
        # Colors
        header_color = (79, 129, 189)
        subheader_color = (220, 230, 241)
        
        def add_header(text):
            pdf.set_font("helvetica", 'B', 14)
            pdf.set_fill_color(*header_color)
            pdf.set_text_color(255, 255, 255)
            pdf.cell(0, 10, text, border=1, ln=1, align='C', fill=True)
            pdf.set_text_color(0, 0, 0) # reset
            
        def add_subheader(text):
            pdf.set_font("helvetica", 'B', 12)
            pdf.set_fill_color(*subheader_color)
            pdf.cell(0, 8, text, border=1, ln=1, fill=True)
            
        def add_row(k, v, w1=95, w2=95):
            pdf.set_font("helvetica", 'B', 10)
            pdf.cell(w1, 7, str(k), border=1)
            pdf.set_font("helvetica", '', 10)
            pdf.cell(w2, 7, str(v) if v is not None else "N/A", border=1, ln=1)

        # Title
        add_header(f"Inspection Report: Box {box.box_number or 'Unnumbered'}")
        pdf.ln(5)

        # 1. Lot Details
        add_subheader("1. Lot Details")
        add_row("Lot ID", lot.id)
        add_row("Lot Number", lot.lot_number)
        add_row("Lot Inspection Status", lot.inspection_status)
        add_row("Box Type", lot.box_type)
        add_row("Total Boxes", lot.number_of_box)
        add_row("Total Fuzes", lot.number_of_fuze)
        add_row("Fuze Type", lot.fuze_type)
        pdf.ln(3)

        # 2. Box Details
        add_subheader("2. Box Details")
        add_row("Box ID", box.id)
        add_row("Box Number", box.box_number)
        add_row("Sequence Number", box.sequence_number)
        add_row("Box Type", box.box_type)
        add_row("Inspection Status", box.box_inspection_status)
        add_row("Acceptance Status", "Passed" if box.acceptance_status is True else ("Rejected" if box.acceptance_status is False else "Pending/Not Answered"))
        add_row("Inspecting User", user_info)
        add_row("Start Time", box.start_time.strftime("%Y-%m-%d %H:%M:%S") if box.start_time else "N/A")
        add_row("End Time", box.end_time.strftime("%Y-%m-%d %H:%M:%S") if box.end_time else "N/A")
        pdf.ln(3)
        
        # 3. Check List
        add_subheader("3. Inspection Checklist")
        if inspection:
            for column in inspection.__table__.columns:
                if column.name not in ["id"]:
                    val = getattr(inspection, column.name)
                    pretty_key = column.name.replace("_", " ").title()
                    status_text = "N/A"
                    if val is True:
                        status_text = "Pass"
                    elif val is False:
                        status_text = "Fail"
                    add_row(pretty_key, status_text)
        else:
            add_row("Checklist Status", "No checklist found")
        pdf.ln(3)
        
        # # 4. Logs
        # add_subheader("4. Box Logs")
        # if box_logs:
        #     pdf.set_font("helvetica", 'B', 10)
        #     pdf.cell(60, 7, "Timestamp", border=1)
        #     pdf.cell(130, 7, "Description", border=1, ln=1)
        #     pdf.set_font("helvetica", '', 10)
        #     for log in box_logs:
        #         pdf.cell(60, 7, log.date_time.strftime("%Y-%m-%d %H:%M:%S") if log.date_time else "N/A", border=1)
        #         pdf.cell(130, 7, log.description, border=1, ln=1)
        # else:
        #     add_row("Status", "No logs found")
        # pdf.ln(5)

        # 5. Attached Media
        
        if box.files_path and os.path.exists(box.files_path):
            has_media = False
            for f in os.listdir(box.files_path):
                if f.lower().endswith(('.png', '.jpg', '.jpeg')):
                    if not has_media:
                        add_subheader("5. Attached Images")
                        pdf.ln(2)
                        has_media = True
                    img_path = os.path.join(box.files_path, f)
                    try:
                        # Add image fitting mostly in width
                        pdf.image(img_path, w=150)
                        pdf.ln(5)
                    except Exception as e:
                        pdf.set_font("helvetica", 'I', 8)
                        pdf.cell(0, 5, f"Could not load image {f} ({str(e)})", ln=1)

        pdf_bytes = pdf.output()
        return io.BytesIO(pdf_bytes)

    @staticmethod
    async def generate_fuze_excel_report(db: Session, fuze_id: int):
        try:
            fuze = db.exec(select(Fuze).where(Fuze.id == fuze_id)).first()
            if not fuze:
                raise HTTPException(status_code=404, detail="Fuze not found")
            
            lot = db.exec(select(Lot).where(Lot.id == fuze.lot_id)).first()
            if not lot:
                raise HTTPException(status_code=404, detail="Associated lot not found")
            
            fuze_logs = db.exec(select(FuzeLog).where(FuzeLog.fuze_id == fuze_id).order_by(FuzeLog.date_time)).all()
            
            # Determine inspection
            inspection = None
            if fuze.primer_1al_inspection_id:
                inspection = db.exec(
                    select(Primer1ALInspection).where(Primer1ALInspection.id == fuze.primer_1al_inspection_id)).first()
                    
            user_info = "N/A"
            if fuze.user_id:
                user = db.exec(select(User).where(User.id == fuze.user_id)).first()
                if user:
                    user_info = f"({user.email})"

            # Initialize workbook
            wb = Workbook()
            ws = wb.active
            ws.title = f"Fuze Sequence {fuze.sequence_number} Report"

            # Call helper
            ReportService._build_fuze_excel_sheet(ws, fuze, lot, user_info, inspection, fuze_logs)
            
            # Save to memory
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            return output
        except Exception as e:
            print(e)
            raise HTTPException(status_code=500, detail=str(e))

    @staticmethod
    async def generate_fuze_pdf_report(db: Session, fuze_id: int):
        try:
            fuze = db.exec(select(Fuze).where(Fuze.id == fuze_id)).first()
            if not fuze:
                raise HTTPException(status_code=404, detail="Fuze not found")
            
            lot = db.exec(select(Lot).where(Lot.id == fuze.lot_id)).first()
            if not lot:
                raise HTTPException(status_code=404, detail="Associated lot not found")
            
            fuze_logs = db.exec(select(FuzeLog).where(FuzeLog.fuze_id == fuze_id).order_by(FuzeLog.date_time)).all()
            
            # Determine inspection
            inspection = None
            if fuze.primer_1al_inspection_id:
                inspection = db.exec(
                    select(Primer1ALInspection).where(Primer1ALInspection.id == fuze.primer_1al_inspection_id)).first()
                    
            user_info = "N/A"
            if fuze.user_id:
                user = db.exec(select(User).where(User.id == fuze.user_id)).first()
                if user:
                    user_info = f"({user.email})"

            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("helvetica", size=10)
            
            # Colors
            header_color = (79, 129, 189)
            subheader_color = (220, 230, 241)
            
            def add_header(text):
                pdf.set_font("helvetica", 'B', 14)
                pdf.set_fill_color(*header_color)
                pdf.set_text_color(255, 255, 255)
                pdf.cell(0, 10, text, border=1, ln=1, align='C', fill=True)
                pdf.set_text_color(0, 0, 0) # reset
                
            def add_subheader(text):
                pdf.set_font("helvetica", 'B', 12)
                pdf.set_fill_color(*subheader_color)
                pdf.cell(0, 8, text, border=1, ln=1, fill=True)
                
            def add_row(k, v, w1=95, w2=95):
                pdf.set_font("helvetica", 'B', 10)
                pdf.cell(w1, 7, str(k), border=1)
                pdf.set_font("helvetica", '', 10)
                pdf.cell(w2, 7, str(v) if v is not None else "N/A", border=1, ln=1)

            # Title
            add_header(f"Inspection Report: Fuze Sequence {fuze.sequence_number}")
            pdf.ln(5)

            # 1. Lot Details
            add_subheader("1. Lot Details")
            add_row("Lot ID", lot.id)
            add_row("Lot Number", lot.lot_number)
            add_row("Lot Inspection Status", lot.inspection_status)
            add_row("Box Type", lot.box_type)
            add_row("Total Boxes", lot.number_of_box)
            add_row("Total Fuzes", lot.number_of_fuze)
            add_row("Fuze Type", lot.fuze_type)
            pdf.ln(3)

            # 2. Fuze Details
            add_subheader("2. Fuze Details")
            add_row("Fuze ID", fuze.id)
            add_row("Sequence Number", fuze.sequence_number)
            add_row("Fuze Type", fuze.fuze_type)
            add_row("Inspection Status", fuze.fuze_inspection_status)
            add_row("Acceptance Status", "Passed" if fuze.acceptance_status is True else ("Rejected" if fuze.acceptance_status is False else "Pending/Not Answered"))
            add_row("Inspecting User", user_info)
            add_row("Start Time", fuze.start_time.strftime("%Y-%m-%d %H:%M:%S") if fuze.start_time else "N/A")
            add_row("End Time", fuze.end_time.strftime("%Y-%m-%d %H:%M:%S") if fuze.end_time else "N/A")
            pdf.ln(3)
            
            # 3. Check List
            add_subheader("3. Inspection Checklist")
            if inspection:
                for column in inspection.__table__.columns:
                    if column.name not in ["id"]:
                        val = getattr(inspection, column.name)
                        pretty_key = column.name.replace("_", " ").title()
                        status_text = "N/A"
                        if val is True:
                            status_text = "Pass"
                        elif val is False:
                            status_text = "Fail"
                        add_row(pretty_key, status_text)
            else:
                add_row("Checklist Status", "No checklist found")
            pdf.ln(3)
            
            # # 4. Logs
            # add_subheader("4. Fuze Logs")
            # if fuze_logs:
            #     pdf.set_font("helvetica", 'B', 10)
            #     pdf.cell(60, 7, "Timestamp", border=1)
            #     pdf.cell(130, 7, "Description", border=1, ln=1)
            #     pdf.set_font("helvetica", '', 10)
            #     for log in fuze_logs:
            #         pdf.cell(60, 7, log.date_time.strftime("%Y-%m-%d %H:%M:%S") if log.date_time else "N/A", border=1)
            #         pdf.cell(130, 7, log.description, border=1, ln=1)
            # else:
            #     add_row("Status", "No logs found")
            # pdf.ln(5)

            # 5. Attached Media
            if fuze.files_path and os.path.exists(fuze.files_path):
                has_media = False
                for f in os.listdir(fuze.files_path):
                    if f.lower().endswith(('.png', '.jpg', '.jpeg')):
                        if not has_media:
                            add_subheader("5. Attached Images")
                            pdf.ln(2)
                            has_media = True
                        img_path = os.path.join(fuze.files_path, f)
                        try:
                            # Add image fitting mostly in width
                            pdf.image(img_path, w=150)
                            pdf.ln(5)
                        except Exception as e:
                            pdf.set_font("helvetica", 'I', 8)
                            pdf.cell(0, 5, f"Could not load image {f} ({str(e)})", ln=1)

            pdf_bytes = pdf.output()
            return io.BytesIO(pdf_bytes)
        except Exception as e:
            print(e)
            raise HTTPException(status_code=500, detail=str(e))

    @staticmethod
    async def generate_lot_export_zip(db: Session, lot_id: int):
        try:
            lot = db.exec(select(Lot).where(Lot.id == lot_id)).first()
            if not lot:
                raise HTTPException(status_code=404, detail="Lot not found")
            
            if lot.inspection_status != "Completed":
                pass # We can still allow export, but perhaps warn. Requirements say usually for completed.

            # Fetch all boxes and fuzes
            boxes = db.exec(select(Box).where(Box.lot_id == lot_id).order_by(Box.sequence_number)).all()
            fuzes = db.exec(select(Fuze).where(Fuze.lot_id == lot_id).order_by(Fuze.sequence_number)).all()

            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
                
                # --- 1. lot_details.xlsx ---
                lot_wb = Workbook()
                ws = lot_wb.active
                ws.title = "Lot Details"
                
                header_font = Font(bold=True, size=14, color="FFFFFF")
                header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
                label_font = Font(bold=True)
                thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

                def add_data(ws, row, label, val):
                    l_cell = ws.cell(row=row, column=1, value=label)
                    l_cell.font = label_font
                    l_cell.border = thin_border
                    v_cell = ws.cell(row=row, column=2, value=str(val) if val is not None else "N/A")
                    v_cell.border = thin_border
                    return row + 1

                ws.cell(row=1, column=1, value=f"Lot {lot.lot_number} Report").font = header_font
                ws.cell(row=1, column=1).fill = header_fill
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
                
                r = 2
                r = add_data(ws, r, "Lot ID", lot.id)
                r = add_data(ws, r, "Lot Number", lot.lot_number)
                r = add_data(ws, r, "Inspection Status", lot.inspection_status)
                r = add_data(ws, r, "Acceptance Status", lot.acceptance_status)
                r = add_data(ws, r, "Total Boxes", lot.number_of_box)
                r = add_data(ws, r, "Total Fuzes", lot.number_of_fuze)
                r = add_data(ws, r, "Box Type", lot.box_type)
                r = add_data(ws, r, "Fuze Type", lot.fuze_type)
                r = add_data(ws, r, "Created At", lot.created_at.strftime("%Y-%m-%d %H:%M:%S") if lot.created_at else "N/A")
                
                ws.column_dimensions['A'].width = 30
                ws.column_dimensions['B'].width = 40
                
                lot_out = io.BytesIO()
                lot_wb.save(lot_out)
                zip_file.writestr(f"Lot_{lot.lot_number}/lot_details.xlsx", lot_out.getvalue())
                
                # --- 2. box_details.xlsx ---
                if boxes:
                    box_wb = Workbook()
                    box_wb.remove(box_wb.active) # Remove default sheet
                    for box in boxes:
                        ws = box_wb.create_sheet(title=f"Box {box.sequence_number}")
                        box_logs = db.exec(select(BoxLog).where(BoxLog.box_id == box.id).order_by(BoxLog.date_time)).all()
                        inspection = None
                        if getattr(box, "box_h2a_inspection_id", None):
                            inspection = db.exec(select(BoxH2AInspection).where(BoxH2AInspection.id == box.box_h2a_inspection_id)).first()
                        elif getattr(box, "box_type", None) in ["H2A", "Box H2A", "H2A Box"]:
                            if getattr(box, "box_h2a_inspection_id", None):
                                inspection = db.exec(select(BoxH2AInspection).where(BoxH2AInspection.id == box.box_h2a_inspection_id)).first()
                        elif getattr(box, "box_type", None) in ["M-104 Type 2", "M104"] and hasattr(box, "box_m104_inspection_id") and box.box_m104_inspection_id:
                            inspection = db.exec(select(BoxM104Inspection).where(BoxM104Inspection.id == box.box_m104_inspection_id)).first()
                        elif hasattr(box, "box_g70_inspection_id") and box.box_g70_inspection_id:
                            inspection = db.exec(select(BoxG70Inspection).where(BoxG70Inspection.id == box.box_g70_inspection_id)).first()
                        user_info = "N/A"
                        if box.user_id:
                            user = db.exec(select(User).where(User.id == box.user_id)).first()
                            if user: user_info = f"({user.email})"
                        
                        ReportService._build_box_excel_sheet(ws, box, lot, user_info, inspection, box_logs)
                        
                    box_out = io.BytesIO()
                    box_wb.save(box_out)
                    zip_file.writestr(f"Lot_{lot.lot_number}/box_details.xlsx", box_out.getvalue())

                # --- 3. fuze_details.xlsx ---
                if fuzes:
                    fuze_wb = Workbook()
                    fuze_wb.remove(fuze_wb.active)
                    for fuze in fuzes:
                        ws = fuze_wb.create_sheet(title=f"Fuze {fuze.sequence_number}")
                        fuze_logs = db.exec(select(FuzeLog).where(FuzeLog.fuze_id == fuze.id).order_by(FuzeLog.date_time)).all()
                        inspection = None
                        if fuze.primer_1al_inspection_id:
                            inspection = db.exec(select(Primer1ALInspection).where(Primer1ALInspection.id == fuze.primer_1al_inspection_id)).first()
                        user_info = "N/A"
                        if fuze.user_id:
                            user = db.exec(select(User).where(User.id == fuze.user_id)).first()
                            if user: user_info = f"({user.email})"
                            
                        ReportService._build_fuze_excel_sheet(ws, fuze, lot, user_info, inspection, fuze_logs)
                        
                    fuze_out = io.BytesIO()
                    fuze_wb.save(fuze_out)
                    zip_file.writestr(f"Lot_{lot.lot_number}/fuze_details.xlsx", fuze_out.getvalue())

                # --- 4. Box Folders (PDF + Media) ---
                for box in boxes:
                    folder_name = f"Lot_{lot.lot_number}/Boxes/Box_{box.sequence_number}"
                    try:
                        pdf_stream = await ReportService.generate_box_pdf_report(db, box.id)
                        zip_file.writestr(f"{folder_name}/report.pdf", pdf_stream.getvalue())
                    except Exception:
                        pass
                    
                    if box.files_path and os.path.exists(box.files_path):
                        for f in os.listdir(box.files_path):
                            if f.lower().endswith(('.png', '.jpg', '.jpeg', '.mp4', '.avi', '.mov')):
                                filepath = os.path.join(box.files_path, f)
                                zip_file.write(filepath, f"{folder_name}/{f}")

                # --- 5. Fuze Folders (PDF + Media) ---
                for fuze in fuzes:
                    folder_name = f"Lot_{lot.lot_number}/Fuzes/Fuze_{fuze.sequence_number}"
                    try:
                        pdf_stream = await ReportService.generate_fuze_pdf_report(db, fuze.id)
                        zip_file.writestr(f"{folder_name}/report.pdf", pdf_stream.getvalue())
                    except Exception:
                        pass
                    
                    if fuze.files_path and os.path.exists(fuze.files_path):
                        for f in os.listdir(fuze.files_path):
                            if f.lower().endswith(('.png', '.jpg', '.jpeg', '.mp4', '.avi', '.mov')):
                                filepath = os.path.join(fuze.files_path, f)
                                zip_file.write(filepath, f"{folder_name}/{f}")

            zip_buffer.seek(0)
            return zip_buffer
        except Exception as e:
            print(e)
            raise HTTPException(status_code=500, detail=str(e))

ReportService.generate_primer_excel_report = ReportService.generate_fuze_excel_report
ReportService.generate_primer_pdf_report = ReportService.generate_fuze_pdf_report

